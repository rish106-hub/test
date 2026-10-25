"""Isolated 03-nav build + LibreOffice recalculation for the self-check tests.

Mirrors build/generate_workbook.py::build for `--test-build --only 03-nav` (interfaces stubbed from the
synthetic fixtures), but lets a test edit the fixture rows and start controls away from default.
"""
from __future__ import annotations

import copy
import datetime as dt
from contextlib import contextmanager
from pathlib import Path

from build import main_sheets
from build.data_loader import load_model_data
from build.generate_workbook import SHEET_ORDER
from build.interface_stubs import add_stubs
from build.recalc import read_values, recalc
from src.common.workbook import WorkbookContext
from src.fund_nav.sheets import build_sheets
from src.scenarios import sheets as scenario_sheets


class Built:
    def __init__(self, ctx: WorkbookContext, values: dict, data):
        self.ctx, self.values, self.data = ctx, values, data

    def table(self, name: str) -> list[dict]:
        """Recalculated rows of an Excel table as {header: value}."""
        info = self.ctx.tables[name]
        sheet = self.values[info["sheet"]]
        from openpyxl.utils import range_boundaries
        c0, r0, c1, r1 = range_boundaries(":".join(info["range"]))
        from openpyxl.utils import get_column_letter as L
        rows = []
        for r in range(r0 + 1, r1 + 1):
            rows.append({h: sheet.get(f"{L(c0 + j)}{r}") for j, h in enumerate(info["columns"])})
        return rows

    def row(self, name: str, key_col: str, key: str) -> dict:
        return next(r for r in self.table(name) if r[key_col] == key)


@contextmanager
def _overrides(controls, sector_shocks):
    old = (dict(scenario_sheets.PRICE_CONTROL_OVERRIDES), dict(scenario_sheets.SECTOR_SHOCK_OVERRIDES))
    scenario_sheets.PRICE_CONTROL_OVERRIDES.update(controls or {})
    scenario_sheets.SECTOR_SHOCK_OVERRIDES.update(sector_shocks or {})
    try:
        yield
    finally:
        scenario_sheets.PRICE_CONTROL_OVERRIDES.clear()
        scenario_sheets.PRICE_CONTROL_OVERRIDES.update(old[0])
        scenario_sheets.SECTOR_SHOCK_OVERRIDES.clear()
        scenario_sheets.SECTOR_SHOCK_OVERRIDES.update(old[1])


def build_isolated(tmp_path: Path, *, edit=None, controls=None, sector_shocks=None, extra_builders=()) -> Built:
    """controls: {(Security_ID, column): value}; sector_shocks: {Sector: shock}; edit(data) mutates fixture rows."""
    data = load_model_data(True)
    data.tables = copy.deepcopy(data.tables)
    if edit:
        edit(data)
    path = tmp_path / "isolated.xlsx"
    ctx = WorkbookContext(str(path), test_build=True)
    ctx.create_sheets(SHEET_ORDER)
    with _overrides(controls, sector_shocks):
        main_sheets.build_sheets(ctx, data)
        build_sheets(ctx, data)
    add_stubs(ctx, data, {"00-main", "03-nav"})
    main_sheets.build_dashboard(ctx, data)
    ctx.close()
    values = read_values(recalc(path, tmp_path / "recalc"))
    return Built(ctx, values, data)


def set_global(data, name: str, value: str) -> None:
    for r in data.tables["global_inputs"]:
        if r["Name"] == name:
            r["Value"] = value
            return
    raise KeyError(name)


def same(expected, actual, rel: float = 1e-9) -> bool:
    """Workbook value vs reference value. None <=> '#N/A' (an unmarkable position)."""
    if expected is None:
        return actual in ("#N/A", None, "")
    if isinstance(expected, dt.date):
        return actual is not None and (actual.date() if isinstance(actual, dt.datetime) else actual) == expected
    if isinstance(expected, bool) or isinstance(expected, str):
        return actual == expected
    if actual in (None, "") or isinstance(actual, str):
        return False
    return abs(float(actual) - float(expected)) <= rel * max(1.0, abs(float(expected)))
