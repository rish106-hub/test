"""LibreOffice recalculation harness.

XlsxWriter stores 0 as the cached result of every formula. Reading such a file
with openpyxl(data_only=True) returns zeros, so a naive check would pass for any
formula. This harness opens the workbook in headless LibreOffice with a profile
that forces a full recalculation on load (build/lo_profile), saves a copy, and
returns its path. Read the copy with openpyxl(data_only=True) to get real values.

usage: python build/recalc.py build/out/Indian_Equity_Valuation_and_MF_NAV_Model.xlsx
"""
from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PROFILE_TEMPLATE = ROOT / "build" / "lo_profile"


def soffice_available() -> bool:
    return shutil.which("soffice") is not None


def recalc(xlsx: str | Path, out_dir: str | Path | None = None, timeout: int = 300) -> Path:
    xlsx = Path(xlsx).resolve()
    out_dir = Path(out_dir or ROOT / "build" / "recalc").resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        profile = Path(tmp) / "profile"
        shutil.copytree(PROFILE_TEMPLATE, profile)
        cmd = ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--calc",
               "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out_dir), str(xlsx)]
        subprocess.run(cmd, check=True, timeout=timeout, capture_output=True)
    result = out_dir / xlsx.name
    if not result.exists():
        raise RuntimeError(f"LibreOffice did not produce {result}")
    return result


def read_values(xlsx: str | Path) -> dict[str, dict[str, object]]:
    """{sheet: {cell: value}} for every non-empty cell of a recalculated workbook."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    return {ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
            for ws in wb.worksheets}


def read_defined_names(xlsx: str | Path) -> dict[str, object]:
    """Value of every single-cell defined name in a recalculated workbook."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    out = {}
    for name, dn in wb.defined_names.items():
        for sheet, ref in dn.destinations:
            out[name] = wb[sheet][ref.replace("$", "")].value
    return out


if __name__ == "__main__":
    print(recalc(sys.argv[1]))
