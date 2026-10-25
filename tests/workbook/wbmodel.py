"""Read a generated workbook for the W and E layers.

Two files are always read together:
  * the generated workbook (formulas, fonts, tables, names) as written by build/generate_workbook.py;
  * its LibreOffice recalculation (build/recalc.py), the only source of values. XlsxWriter caches 0
    for every formula, so cached values are never trusted.
"""
from __future__ import annotations

import json
import math
import re
import shutil
from dataclasses import dataclass, field
from datetime import date, datetime
from functools import cached_property
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

from src.common import excel_contract as X
from src.common import ownership

ERRORS = ("#REF!", "#NAME?", "#DIV/0!", "#VALUE!", "#N/A", "#NUM!", "#NULL!", "Err:")


def is_error(v) -> bool:
    return isinstance(v, str) and v.startswith(ERRORS)


def blankish(v) -> bool:
    return v is None or v == "" or (isinstance(v, str) and v.startswith("#N/A"))


def num(v):
    """Workbook value -> float, or None for blank / text / error."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


def as_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 20000 < v < 80000:   # serial date
        return date(1899, 12, 30).fromordinal(date(1899, 12, 30).toordinal() + int(v))
    if isinstance(v, str):
        try:
            return date.fromisoformat(v[:10])
        except ValueError:
            return None
    return None


# --------------------------------------------------------------------------- formula references
_STRING = re.compile(r'"(?:[^"]|"")*"')
_SHEET_REF = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z_][A-Za-z0-9_.]*))!(\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)")
_LOCAL_REF = re.compile(r"(?<![A-Za-z0-9_.!\]$'#])(\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)(?![A-Za-z0-9_(\[])")


@dataclass
class Refs:
    columns: set = field(default_factory=set)      # (table, column, this_row: bool)
    names: set = field(default_factory=set)
    cells: set = field(default_factory=set)        # (sheet or None for same sheet, range text)


def _match_bracket(s: str, i: int) -> int:
    depth = 0
    for j in range(i, len(s)):
        if s[j] == "[":
            depth += 1
        elif s[j] == "]":
            depth -= 1
            if depth == 0:
                return j
    return len(s) - 1


def parse_refs(formula: str) -> Refs:
    """Structured, named and A1 references of one formula (strings stripped first)."""
    f = _STRING.sub('""', formula or "")
    refs = Refs()
    out, i = [], 0
    for m in re.finditer(r"(tbl_[A-Za-z0-9_]+)?\[", f):
        if m.start() < i:
            continue
        table = m.group(1)
        j = _match_bracket(f, m.end() - 1)
        inner = f[m.end() - 1: j + 1]
        out.append(f[i:m.start()])
        i = j + 1
        items = re.findall(r"\[([^\[\]]+)\]", inner) or [inner.strip("[]")]
        this_row = any(x.strip().lower() in ("#this row", "@") for x in items) or inner.startswith("[@")
        for it in items:
            it = it.strip().lstrip("@")
            if it and not it.startswith("#"):
                refs.columns.add((table, it, this_row))
    out.append(f[i:])
    rest = " ".join(out)
    refs.names = set(X.names_referenced(rest))
    for m in _SHEET_REF.finditer(rest):
        sheet = (m.group(1) or m.group(2)).replace("''", "'")
        refs.cells.add((sheet, m.group(3).replace("$", "")))
    rest2 = _SHEET_REF.sub(" ", rest)
    rest2 = re.sub(r"\b(?:in|out)_[A-Za-z0-9_]+", " ", rest2)
    for m in _LOCAL_REF.finditer(rest2):
        refs.cells.add((None, m.group(1).replace("$", "")))
    return refs


# --------------------------------------------------------------------------- workbook model
@dataclass
class Table:
    name: str
    sheet: str
    ref: str
    headers: list
    min_col: int
    header_row: int
    max_row: int

    def col_letter(self, header: str) -> str:
        return get_column_letter(self.min_col + self.headers.index(header))

    def cell(self, header: str, row: int) -> str:
        return f"{self.col_letter(header)}{row}"

    @property
    def data_rows(self) -> range:
        return range(self.header_row + 1, self.max_row + 1)


class Workbook:
    def __init__(self, source: Path, recalculated: Path, build_dir: Path | None = None):
        self.source, self.recalculated = Path(source), Path(recalculated)
        self.build_dir = Path(build_dir) if build_dir else self.source.parent
        self.wf = openpyxl.load_workbook(self.source)                      # formulas + fonts
        self.wv = openpyxl.load_workbook(self.recalculated, data_only=True)  # recalculated values

    # ---------------------------------------------------------------- build metadata
    @cached_property
    def manifest(self) -> dict:
        p = self.build_dir / "manifest.json"
        return json.loads(p.read_text()) if p.exists() else {}

    @cached_property
    def build_report(self) -> dict:
        p = self.build_dir / "build_report.json"
        return json.loads(p.read_text()) if p.exists() else {}

    @property
    def is_test_build(self) -> bool:
        return bool(self.build_report.get("test_build", self.manifest.get("test_build", True)))

    def table_owner(self, name: str) -> str | None:
        t = self.manifest.get("tables", {}).get(name)
        return t["owner"] if t else (ownership.sheet_owner(self.tables[name].sheet) if name in self.tables else None)

    def is_real(self, name: str) -> bool:
        """Table built by its producing module (not an interface stub, not absent)."""
        return name in self.tables and self.table_owner(name) not in (None, "STUB")

    # ---------------------------------------------------------------- tables and names
    @cached_property
    def tables(self) -> dict[str, Table]:
        out = {}
        for ws in self.wf.worksheets:
            for name, t in ws.tables.items():
                ref = t.ref if hasattr(t, "ref") else t
                c0, r0, c1, r1 = range_boundaries(ref)
                headers = [ws.cell(r0, c).value for c in range(c0, c1 + 1)]
                out[name] = Table(name, ws.title, ref, [str(h) for h in headers], c0, r0, r1)
        return out

    @cached_property
    def names(self) -> dict[str, tuple[str, str]]:
        out = {}
        for name, dn in self.wf.defined_names.items():
            for sheet, ref in dn.destinations:
                out[name] = (sheet, ref.replace("$", ""))
        return out

    def value(self, sheet: str, cell: str):
        return self.wv[sheet][cell].value

    def formula(self, sheet: str, cell: str):
        v = self.wf[sheet][cell].value
        return v if isinstance(v, str) and v.startswith("=") else None

    def name_value(self, name: str):
        s, c = self.names[name]
        return self.value(s, c)

    def rows(self, table: str, values: dict | None = None) -> list[dict]:
        """Table rows as {header: value} from the recalculated file (or from a values dump)."""
        t = self.tables[table]
        out = []
        for r in t.data_rows:
            if values is None:
                row = {h: self.value(t.sheet, t.cell(h, r)) for h in t.headers}
            else:
                row = {h: values.get(t.sheet, {}).get(t.cell(h, r)) for h in t.headers}
            if any(v not in (None, "") for v in row.values()):
                row["__row__"] = r
                out.append(row)
        return out

    def row_index(self, table: str, key_col: str, key) -> int | None:
        t = self.tables[table]
        for r in t.data_rows:
            if self.wf[t.sheet][t.cell(key_col, r)].value == key or self.value(t.sheet, t.cell(key_col, r)) == key:
                return r
        return None

    # ---------------------------------------------------------------- formulas
    def formula_cells(self):
        """(sheet, coordinate, formula) for every formula cell of the generated workbook."""
        for ws in self.wf.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        yield ws.title, c.coordinate, c.value

    def table_of(self, sheet: str, coord: str) -> tuple[Table, str] | None:
        col, row = openpyxl.utils.cell.coordinate_from_string(coord)
        ci = column_index_from_string(col)
        for t in self.tables.values():
            if t.sheet == sheet and t.header_row < row <= t.max_row and t.min_col <= ci < t.min_col + len(t.headers):
                return t, t.headers[ci - t.min_col]
        return None

    def precedents(self, sheet: str, coord: str, formula: str) -> set[tuple[str, str]]:
        """Direct precedent cells of one formula cell."""
        refs = parse_refs(formula)
        _, row = openpyxl.utils.cell.coordinate_from_string(coord)
        here = self.table_of(sheet, coord)
        out = set()
        for table, col, this_row in refs.columns:
            t = self.tables.get(table) if table else (here[0] if here else None)
            if t is None or col not in t.headers:
                continue
            rows = [row] if this_row else list(t.data_rows)
            out |= {(t.sheet, t.cell(col, r)) for r in rows}
        for n in refs.names:
            if n in self.names:
                out.add(self.names[n])
        for s, rng in refs.cells:
            s = s or sheet
            c0, r0, c1, r1 = range_boundaries(rng if ":" in rng else f"{rng}:{rng}")
            if (c1 - c0 + 1) * (r1 - r0 + 1) > 20000:
                continue
            out |= {(s, f"{get_column_letter(c)}{r}") for c in range(c0, c1 + 1) for r in range(r0, r1 + 1)}
        return out

    @cached_property
    def graph(self) -> dict[tuple[str, str], set]:
        return {(s, c): self.precedents(s, c, f) for s, c, f in self.formula_cells()}

    def upstream(self, cells) -> set:
        """Transitive precedents of a set of cells."""
        seen, stack = set(), list(cells)
        while stack:
            node = stack.pop()
            for p in self.graph.get(node, ()):
                if p not in seen:
                    seen.add(p)
                    stack.append(p)
        return seen

    def column_cells(self, table: str, col: str) -> set:
        t = self.tables[table]
        return {(t.sheet, t.cell(col, r)) for r in t.data_rows}

    def font_rgb(self, sheet: str, coord: str) -> str | None:
        c = self.wf[sheet][coord]
        col = c.font.color if c.font else None
        rgb = getattr(col, "rgb", None) if col is not None else None
        return rgb[-6:].upper() if isinstance(rgb, str) else None

    def all_values(self) -> dict[str, dict[str, object]]:
        return {ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
                for ws in self.wv.worksheets}


# --------------------------------------------------------------------------- perturbation harness
def recalc_values(xlsx: Path, out_dir: Path) -> dict[str, dict[str, object]]:
    from build.recalc import read_values, recalc
    return read_values(recalc(xlsx, out_dir))


def perturbed_values(wb: Workbook, edits: list[tuple[str, str, object]], work_dir: Path) -> dict:
    """Copy the generated workbook, set input cells with openpyxl, recalculate, return all values.

    openpyxl drops charts but keeps formulas, tables and defined names; test_W_10_00 proves an
    unperturbed round trip recalculates to identical values before any perturbation is trusted.
    """
    work_dir = Path(work_dir)
    work_dir.mkdir(parents=True, exist_ok=True)
    src = openpyxl.load_workbook(wb.source)
    for sheet, coord, value in edits:
        src[sheet][coord].value = value
    path = work_dir / wb.source.name
    src.save(path)
    return recalc_values(path, work_dir / "recalc")


VOLATILE = re.compile(r"\b(TODAY|NOW|RAND|RANDBETWEEN)\s*\(")


def value_diff(a: dict, b: dict, *, ignore=frozenset(), rel: float = 1e-12) -> dict:
    """{(sheet, cell): (a, b)} for every cell whose value differs."""
    out = {}
    for sheet in set(a) | set(b):
        sa, sb = a.get(sheet, {}), b.get(sheet, {})
        for cell in set(sa) | set(sb):
            if (sheet, cell) in ignore:
                continue
            va, vb = sa.get(cell), sb.get(cell)
            if va == vb:
                continue
            if isinstance(va, (int, float)) and isinstance(vb, (int, float)) and not isinstance(va, bool) \
                    and math.isclose(va, vb, rel_tol=rel, abs_tol=1e-15):
                continue
            out[(sheet, cell)] = (va, vb)
    return out


def volatile_cells(wb: Workbook) -> frozenset:
    return frozenset((s, c) for s, c, f in wb.formula_cells() if VOLATILE.search(f))


def copy_build(src_dir: Path, dst_dir: Path):
    dst_dir.mkdir(parents=True, exist_ok=True)
    for p in Path(src_dir).iterdir():
        if p.is_file():
            shutil.copy2(p, dst_dir / p.name)
