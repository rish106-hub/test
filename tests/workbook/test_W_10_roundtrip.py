"""W-10 prerequisite: the perturbation harness is trustworthy.

Before any spec manual test is run by perturbing input cells with openpyxl, prove that an
unperturbed openpyxl round trip (load, save, recalculate) gives exactly the values of a direct
recalculation, and that formulas, tables and defined names survive it (spec/test_plan.md W-10).
"""
import openpyxl
import pytest
import xlsxwriter

from tests.workbook.wbmodel import Workbook, perturbed_values, recalc_values, value_diff, volatile_cells

pytestmark = pytest.mark.workbook


def test_W_10_00_unperturbed_roundtrip_is_value_identical(built, tmp_path, record_property):
    direct = built.all_values()
    trip = perturbed_values(built, [], tmp_path)
    d = value_diff(direct, trip, ignore=volatile_cells(built))
    n = sum(len(v) for v in direct.values())
    record_property("qa_evidence", f"{n} recalculated cells compared after an openpyxl load/save round trip: "
                                   f"{len(d)} differ")
    assert d == {}, list(d.items())[:10]
    saved = openpyxl.load_workbook(tmp_path / built.source.name)
    assert {t for ws in saved.worksheets for t in ws.tables} == set(built.tables)
    assert set(saved.defined_names) == set(built.names)
    n_formulas = sum(1 for _ in built.formula_cells())
    n_saved = sum(1 for ws in saved.worksheets for row in ws.iter_rows() for c in row
                  if isinstance(c.value, str) and c.value.startswith("="))
    assert n_saved == n_formulas


def test_W_10_01_harness_detects_a_change_through_tables_and_names(tmp_path):
    """Self-test on a small workbook with the features the model uses: a table with this-row
    structured references, an INDEX/MATCH link across sheets and a defined input name."""
    path = tmp_path / "harness.xlsx"
    w = xlsxwriter.Workbook(str(path))
    a, b = w.add_worksheet("IN"), w.add_worksheet("OUT")
    a.write("A1", "shock"); a.write("B1", -0.10)
    w.define_name("in_Shock", "=IN!$B$1")
    a.add_table("A3:C5", {"name": "tbl_P", "columns": [
        {"header": "ID"}, {"header": "P0"},
        {"header": "P1", "formula": "=tbl_P[[#This Row],[P0]]*(1+in_Shock)"}],
        "data": [["S1", 750], ["S2", 100]]})
    b.write_formula("A1", '=INDEX(tbl_P[P1],MATCH("S1",tbl_P[ID],0))')
    b.write_formula("A2", '=SUM(tbl_P[P1])')
    w.close()
    rec = tmp_path / "direct"
    direct = recalc_values(path, rec)
    assert direct["OUT"]["A1"] == pytest.approx(675.0) and direct["OUT"]["A2"] == pytest.approx(765.0)
    wb = Workbook(path, rec / path.name, tmp_path)
    assert perturbed_values(wb, [], tmp_path / "rt")["OUT"] == direct["OUT"]
    after = perturbed_values(wb, [("IN", "B1", -0.20)], tmp_path / "p")
    d = value_diff(direct, after)
    assert after["OUT"]["A1"] == pytest.approx(600.0)
    assert set(d) == {("IN", "B1"), ("IN", "C4"), ("IN", "C5"), ("OUT", "A1"), ("OUT", "A2")}
    assert wb.upstream({("OUT", "A1")}) >= {("IN", "B1"), ("IN", "B4"), ("IN", "C4")}
