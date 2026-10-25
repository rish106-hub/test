"""Excel equals Python for the 02-valuation sheets (D-020, test plan W-03 for out_*).

Builds `--test-build --only 02-valuation` (S90001 on VAL_HDFC_BANK with stubbed interfaces),
recalculates it in LibreOffice and compares every out_* name, the intermediate cells of every
case, the equity bridge, freshness, sensitivity grids, history and tbl_ValuationOutputs with the
Python reference (src/valuations/bank_model.py and history.py) computed from the same inputs, to
1e-9 relative. XIRR is compared to 1e-7 absolute because LibreOffice iterates it to its own tolerance.

The perturbation tests edit input cells in a copy of the built workbook (openpyxl), recalculate,
and check that the outputs move exactly as the reference says. They prove the outputs are
formulas over the inputs, not values written by the generator.
"""
from __future__ import annotations

import copy
import json
import shutil
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest

from build.data_loader import load_model_data
from build.generate_workbook import OUTPUT_NAME, build
from build.recalc import read_defined_names, read_values, recalc
from src.common import schema as S
from src.valuations import history as H
from src.valuations import parameters as P
from src.valuations.bank_model import NOT_REINVESTED, REINVESTED, run_bank_model
from src.valuations.bank_sheet import NA_G_GE_R
from src.valuations.fixtures import augment_test_fixtures
from src.valuations.inputs import gather_bank_spec, model_inputs
from src.valuations.reference import valued_row

pytestmark = [pytest.mark.workbook, pytest.mark.owner("02-valuation"),
              pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")]

SID = "S90001"
REL = 1e-9


# ---------------------------------------------------------------------------- helpers
def blank(v) -> bool:
    return v is None or v == ""


def same(x, y, rel=REL, abs_=1e-12) -> bool:
    """Workbook value x vs reference value y."""
    if y is None:
        return blank(x)
    if isinstance(y, bool):
        return x is y or x == y
    if isinstance(y, date) and not isinstance(y, datetime):
        return isinstance(x, datetime) and x.date() == y
    if isinstance(y, (int, float)):
        return isinstance(x, (int, float)) and not isinstance(x, bool) and x == pytest.approx(y, rel=rel, abs=abs_)
    return x == y


def assert_same(label, x, y, **kw):
    assert same(x, y, **kw), f"{label}: workbook {x!r} != reference {y!r}"


def reference_for(data):
    reg = {r["Security_ID"]: r for r in data.rows("security_registry")}
    spec = gather_bank_spec(data, reg[SID])
    return spec


class Built:
    def __init__(self, directory: Path):
        self.dir = directory
        self.xlsx = directory / OUTPUT_NAME
        self.cell_map = json.loads((directory / "valuation_cell_map.json").read_text())[SID]

    def recalc(self, xlsx: Path | None = None):
        out = recalc(xlsx or self.xlsx, (xlsx or self.xlsx).parent / "recalc")
        return read_values(out), read_defined_names(out)

    def cell(self, key: str) -> str:
        return self.cell_map["cells"][key]


@pytest.fixture(scope="module")
def built(tmp_path_factory) -> Built:
    out = tmp_path_factory.mktemp("val_build")
    report = build(test_build=True, only="02-valuation", out_dir=out)
    assert report["not_implemented"] == {}
    return Built(out)


@pytest.fixture(scope="module")
def recalculated(built):
    return built.recalc()


@pytest.fixture(scope="module")
def data():
    d = load_model_data(test_build=True)
    augment_test_fixtures(d)
    return d


@pytest.fixture(scope="module")
def spec(data):
    return reference_for(data)


@pytest.fixture(scope="module")
def result(spec):
    return run_bank_model(model_inputs(spec))


# ---------------------------------------------------------------------------- out_* names
def test_every_out_name_equals_the_reference(recalculated, result):
    _, names = recalculated
    outs = {k: v for k, v in names.items() if k.startswith("out_")}
    assert set(outs) == {f"out_{SID}_FV_Base", f"out_{SID}_FV_Bear", f"out_{SID}_FV_Bull", f"out_{SID}_TV_Share"}
    for case in P.CASES:
        assert result.fair_values[case] is not None
        assert_same(f"FV_{case}", outs[f"out_{SID}_FV_{case}"], result.fair_values[case])
    assert_same("TV_Share", outs[f"out_{SID}_TV_Share"], result.tv_share_base)
    # Cases are full alternative driver sets (D-008), not Base x (1 +/- k)
    fv = result.fair_values
    assert fv["Bear"] < fv["Base"] < fv["Bull"]
    assert fv["Bear"] / fv["Base"] != pytest.approx(2 - fv["Bull"] / fv["Base"], rel=1e-6)


# ---------------------------------------------------------------------------- intermediate cells
def test_case_blocks_equal_the_reference(built, recalculated, result):
    vals = recalculated[0][built.cell_map["sheet"]]
    v = lambda key: vals.get(built.cell(key))  # noqa: E731
    fields = {"adv": "advances", "nii": "nii", "nonii": "non_interest_income", "opex": "opex", "ppop": "ppop",
              "prov": "provisions", "pat": "pat", "eps": "eps", "dps": "dps", "eq": "equity", "bvps": "bvps",
              "roe": "roe", "iea": "iea", "sh": "diluted_shares"}
    for case in P.CASES:
        cr = result.cases[case]
        for t, py in enumerate(cr.projection, start=1):
            for key, attr in fields.items():
                assert_same(f"{case} {key} t{t}", v(f"{case}.p.{key}.{t}"), getattr(py, attr))
        for t, bv in enumerate(cr.bv_path):
            assert_same(f"{case} BV_{t}", v(f"{case}.bv.{t}"), bv)
        for t, ri in enumerate(cr.ri.residual_incomes, start=1):
            assert_same(f"{case} RI_{t}", v(f"{case}.ri.{t}"), ri)
        checks = {"sum_pvri": cr.ri.pv_residual_income, "ri_tv": cr.ri.terminal_value,
                  "ri_pvtv": cr.ri.pv_terminal_value, "ri_va": cr.ri.value_at_anchor,
                  "ri_share": cr.ri.terminal_value_share, "ddm_pvd": cr.ddm.pv_dividends,
                  "ddm_tv": cr.ddm.terminal_value, "ddm_va": cr.ddm.value_at_anchor,
                  "ddm_share": cr.ddm.terminal_value_share, "jpb_mult": cr.jpb_multiple, "jpb_va": cr.jpb_value_at_anchor,
                  "crgap_max": cr.capital_retention_gap, "cr_status": cr.capital_retention_status,
                  "tv_status": cr.terminal_status, "tflag": cr.terminal_flag, "fv": cr.fair_value}
        for m, x in cr.value_at_valuation.items():
            checks[f"v.{m}"] = x
        er = cr.expected_return
        checks.update({"bvT": er.bvps_at_exit, "exit": er.exit_value, "er_status": er.status, "reinv": er.reinvested})
        for t, x in enumerate(er.dividends_received, start=1):
            checks[f"rec.{t}"] = x
        for t, (d, f) in enumerate(zip(er.cashflow_dates, er.cashflows)):
            checks[f"cfd.{t}"], checks[f"cf.{t}"] = d, f
        for key, x in checks.items():
            assert_same(f"{case} {key}", v(f"{case}.{key}"), x)
        assert_same(f"{case} XIRR", v(f"{case}.xirr"), er.xirr, rel=0, abs_=1e-7)
        assert_same(f"{case} selected ER", v(f"{case}.er"), er.selected, rel=0, abs_=1e-7)


def test_bridge_freshness_rollforward_and_market_equal_the_reference(built, recalculated, result, spec):
    vals = recalculated[0][built.cell_map["sheet"]]
    v = lambda key: vals.get(built.cell(key))  # noqa: E731
    b = result.bridge
    checks = {
        "A": spec.anchor_date, "V": spec.valuation_date, "R": result.latest_results_period_end,
        "price": spec.price, "re": result.cost_of_equity, "bv0": result.bv0,
        "br_bv0": b.bv_reported_at_anchor, "br_rep": b.eps_reported_after_anchor, "br_stub": b.eps_forecast_stub,
        "br_div": b.dividends_paid, "bv_v": b.bv_estimated_at_valuation, "dev": b.interim_eps_deviation_pct,
        "fflag": b.forecast_update_flag, "bridge_status": result.bridge_status,
        "fin_age": result.financials_age_days, "anchor_age": result.anchor_age_days, "fresh": result.freshness_status,
        "roll": result.roll_factor, "divfv": result.dividends_compounded,
        "int_count": result.interim_count, "int_expected": result.interim_expected,
        "int_status": result.interim_coverage_status, "share_status": result.share_action_status,
        "w_status": result.weights_status, "peer_median": spec.peer_pb_median, "peer_n": spec.peer_n_valid,
        "peer_flag": result.peer.flag, "peer_v": result.peer.value, "out.coverage": result.coverage_status,
        "out.tflag": result.terminal_flag_base, "p_status": result.probability_status,
        "p_er": result.probability_weighted_return,
        "m.upside": result.market["upside_base"], "m.mcap": result.market["market_cap_inr_cr"],
        "m.pb_A": result.market["pb_on_bv_at_anchor"], "m.pb_V": result.market["pb_on_bv_at_valuation_estimate"],
        "m.ipb_V": result.market["implied_pb_at_valuation_estimate"],
        "m.pe_trailing": result.market["trailing_pe_on_anchor_fy_eps"], "m.pe_fwd": result.market["forward_pe_model_fy1"],
        "note": valued_row(spec, result)["Data_Confidence_Note"],
    }
    for key, x in checks.items():
        assert_same(key, v(key), x)
    # The reported book value stays captioned with its own date; the bridged value is an estimate (W-19)
    assert v("bv0_caption") == f"Reported BVPS as of {spec.anchor_date.isoformat()}"
    assert v("bv_v_caption") == f"Estimated book value per share at {spec.valuation_date.isoformat()} (estimate)"
    assert v("missing") == 0


def test_sensitivity_grids_equal_the_reference_and_never_show_g_at_or_above_r(built, recalculated, result):
    vals = recalculated[0][built.cell_map["sheet"]]
    v = lambda key: vals.get(built.cell(key))  # noqa: E731
    for name, grid, cols in (("rg", result.grid_rate_growth, result.grid_growths),
                             ("rroe", result.grid_rate_roe, result.grid_roes)):
        for i, rate in enumerate(result.grid_rates):
            assert_same(f"{name} r{i}", v(f"grid.{name}.row.{i}"), rate)
            for j, col in enumerate(cols):
                assert_same(f"{name} col{j}", v(f"grid.{name}.col.{j}"), col)
                x, y = v(f"grid.{name}.{i}.{j}"), grid[i][j]
                if y is None:
                    assert x == NA_G_GE_R, f"{name}[{i}][{j}] should read {NA_G_GE_R}, got {x!r}"
                else:
                    assert_same(f"{name}[{i}][{j}]", x, y)
    # the centre cell of the r x g grid is the Base RI value at V
    assert result.grid_rate_growth[2][2] == pytest.approx(result.cases["Base"].value_at_valuation["RI"], rel=1e-12)


def test_history_block_equals_the_reference(built, recalculated, spec):
    sheet = built.cell_map["sheet"]
    vals = recalculated[0][sheet]
    ref = H.history_values(spec)
    wb = openpyxl.load_workbook(built.xlsx)          # formulas, to find each history cell
    ws = wb[sheet]
    n_checked = n_blank = 0
    for block, periods, metrics, calcs in (("FY", spec.history_fy, H.HISTORY_FY_METRICS, H.FY_CALCS),
                                           ("Q", spec.history_q, H.HISTORY_Q_METRICS, H.Q_CALCS)):
        keys = list(metrics) + [c.key for c in calcs]
        for metric in keys:
            label = H.metric_label(metric) + f" [{metric}]" if metric in H.UNITS else None
            row = next(r for r in range(1, ws.max_row + 1)
                       if isinstance(ws.cell(r, 1).value, str)
                       and (ws.cell(r, 1).value == label or ws.cell(r, 1).value.endswith(f"[{metric}]")
                            or (label is None and ws.cell(r, 1).value.startswith(_calc_label(calcs, metric))))
                       and _in_block(ws, r, block))
            for j, (lab, _) in enumerate(periods):
                x = vals.get(ws.cell(row, 3 + j).coordinate)
                y = ref[block][(metric, lab)]
                assert_same(f"{block} {metric} {lab}", x, y)
                n_checked += 1
                n_blank += y is None
    assert n_checked > 400 and 0 < n_blank < n_checked     # quarterly NIM is missing in the fixture: shown blank


def _calc_label(calcs, key):
    return next(c.label for c in calcs if c.key == key)


def _in_block(ws, row, block):
    """True if `row` lies in the FY (annual) or Q (quarterly) history section."""
    for r in range(row, 0, -1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.startswith("Annual history"):
            return block == "FY"
        if isinstance(v, str) and v.startswith("Quarterly history"):
            return block == "Q"
    return False


# ---------------------------------------------------------------------------- interface tables
def _table_rows(xlsx: Path, values, sheet: str, table: str) -> list[dict]:
    wb = openpyxl.load_workbook(xlsx)
    ws = wb[sheet]
    ref = ws.tables[table].ref
    first, last = ref.split(":")
    min_col, min_row, max_col, max_row = openpyxl.utils.cell.range_boundaries(ref)
    heads = [ws.cell(min_row, c).value for c in range(min_col, max_col + 1)]
    rows = []
    for r in range(min_row + 1, max_row + 1):
        rows.append({h: values[sheet].get(ws.cell(r, c).coordinate) for h, c in zip(heads, range(min_col, max_col + 1))})
    return rows


def test_valuation_outputs_rows_equal_the_reference_and_unvalued_stay_blank(built, recalculated, spec, result, data):
    rows = _table_rows(built.xlsx, recalculated[0], "VALUATION_OUTPUTS", "tbl_ValuationOutputs")
    assert [r["Security_ID"] for r in rows] == [r["Security_ID"] for r in data.rows("security_registry")]
    assert list(rows[0]) == [c.name for c in S.VALUATION_OUTPUTS.columns]
    ref = valued_row(spec, result)
    valued = next(r for r in rows if r["Security_ID"] == SID)
    for col, y in ref.items():
        assert_same(f"tbl_ValuationOutputs[{col}]", valued[col], y)
    for r in rows:
        if r["Security_ID"] == SID:
            continue
        assert r["Valuation_Coverage_Status"] == "UNVALUED"
        for col in ("FV_Base_INR", "FV_Bear_INR", "FV_Bull_INR", "Terminal_Value_Share_Base"):
            assert blank(r[col]), f"{r['Security_ID']} {col} must be blank, never 0 (R009): {r[col]!r}"
        assert r["Data_Confidence_Note"].startswith("UNVALUED: ")


def test_reference_csv_matches_the_workbook(built, recalculated):
    import csv
    rows = _table_rows(built.xlsx, recalculated[0], "VALUATION_OUTPUTS", "tbl_ValuationOutputs")
    with (built.dir / "reference_valuation_outputs.csv").open() as fh:
        ref = list(csv.DictReader(fh))
    assert [r["Security_ID"] for r in ref] == [r["Security_ID"] for r in rows]
    for x, y in zip(rows, ref):
        for col in ("FV_Base_INR", "FV_Bear_INR", "FV_Bull_INR", "Terminal_Value_Share_Base"):
            if y[col] == "":
                assert blank(x[col])
            else:
                assert x[col] == pytest.approx(float(y[col]), rel=REL)
        assert x["Valuation_Coverage_Status"] == y["Valuation_Coverage_Status"]


def test_historical_metrics_table_links_history_and_base_forecasts(built, recalculated, spec, result):
    rows = _table_rows(built.xlsx, recalculated[0], "HISTORICAL_METRICS", "tbl_HistoricalMetrics")
    assert list(rows[0]) == [c.name for c in S.HISTORICAL_METRICS.columns]
    keys = [(r["Entity_ID"], r["Metric_ID"], r["Fiscal_Label"]) for r in rows]
    assert len(keys) == len(set(keys)), "primary key Entity_ID + Metric_ID + Fiscal_Label must be unique"
    fc = [r for r in rows if r["Is_Forecast"]]
    assert {r["Kind"] for r in fc} == {"Modelled"}
    assert {r["Kind"] for r in rows if not r["Is_Forecast"]} == {"Reported", "Calculated"}
    base = result.cases["Base"].projection
    eps = {r["Fiscal_Label"]: r["Value"] for r in fc if r["Metric_ID"] == "EPS_DILUTED"}
    assert list(eps) == list(spec.fy_labels)
    for label, py in zip(spec.fy_labels, base):
        assert eps[label] == pytest.approx(py.eps, rel=REL)
    ref = H.history_values(spec)
    for r in rows:
        if not r["Is_Forecast"]:
            block = "FY" if r["Fiscal_Label"].startswith("FY") else "Q"
            assert_same(f"{r['Metric_ID']} {r['Fiscal_Label']}", r["Value"], ref[block][(r["Metric_ID"], r["Fiscal_Label"])])


def test_assumptions_table_holds_every_row_with_rationale_and_source_type(built, recalculated, data):
    rows = _table_rows(built.xlsx, recalculated[0], "VALUATION_ASSUMPTIONS", "tbl_Assumptions")
    assert len(rows) == len(data.rows("assumptions"))
    for r in rows:
        assert r["Rationale"] and r["Assumption_Source_Type"] and r["Value"] is not None


# ---------------------------------------------------------------------------- perturbation (W-10 style)
def _edit(src: Path, dst: Path, edits: dict[tuple[str, str], object]) -> Path:
    wb = openpyxl.load_workbook(src)
    for (sheet, cell), value in edits.items():
        wb[sheet][cell] = value
    wb.save(dst)
    return dst


def _assumption_cell(xlsx: Path, assumption_id: str) -> str:
    wb = openpyxl.load_workbook(xlsx)
    ws = wb["VALUATION_ASSUMPTIONS"]
    min_col, min_row, max_col, max_row = openpyxl.utils.cell.range_boundaries(ws.tables["tbl_Assumptions"].ref)
    heads = [ws.cell(min_row, c).value for c in range(min_col, max_col + 1)]
    vcol = min_col + heads.index("Value")
    for r in range(min_row + 1, max_row + 1):
        if ws.cell(r, min_col).value == assumption_id:
            return ws.cell(r, vcol).coordinate
    raise KeyError(assumption_id)


def _rerun(data, changes: dict[str, float | None], **inputs):
    d = copy.deepcopy(data)
    rows = []
    for r in d.rows("assumptions"):
        if r["Assumption_ID"] in changes:
            r = dict(r, Value="" if changes[r["Assumption_ID"]] is None else repr(changes[r["Assumption_ID"]]))
        rows.append(r)
    d.tables["assumptions"] = rows
    spec = reference_for(d)
    return spec, run_bank_model(model_inputs(spec, **inputs))


def test_unperturbed_round_trip_recalculates_identically(built, recalculated, tmp_path):
    copy_ = _edit(built.xlsx, tmp_path / OUTPUT_NAME, {})
    _, names = built.recalc(copy_)
    for k, v in recalculated[1].items():
        if k.startswith("out_"):
            assert names[k] == v


def test_base_driver_change_moves_only_the_base_value(built, recalculated, data, tmp_path):
    aid = f"A-{SID}-Base-NIM-ALL"
    new = 0.042
    xlsx = _edit(built.xlsx, tmp_path / OUTPUT_NAME, {("VALUATION_ASSUMPTIONS", _assumption_cell(built.xlsx, aid)): new})
    _, names = built.recalc(xlsx)
    _, res = _rerun(data, {aid: new})
    for case in P.CASES:
        assert_same(f"FV_{case}", names[f"out_{SID}_FV_{case}"], res.fair_values[case])
    before = recalculated[1]
    assert names[f"out_{SID}_FV_Base"] > before[f"out_{SID}_FV_Base"]
    assert names[f"out_{SID}_FV_Bear"] == before[f"out_{SID}_FV_Bear"]      # cases are independent
    assert names[f"out_{SID}_FV_Bull"] == before[f"out_{SID}_FV_Bull"]


def test_missing_input_leaves_fair_value_blank_and_unvalued(built, data, tmp_path):
    aid = f"A-{SID}-Base-TERMINAL_ROE-ALL"
    xlsx = _edit(built.xlsx, tmp_path / OUTPUT_NAME, {("VALUATION_ASSUMPTIONS", _assumption_cell(built.xlsx, aid)): None})
    values, names = built.recalc(xlsx)
    _, res = _rerun(data, {aid: None})
    assert res.fair_values["Base"] is None and res.coverage_status == "UNVALUED"
    assert blank(names[f"out_{SID}_FV_Base"]), "a missing input must give a blank, never 0 (R001, R009)"
    assert_same("FV_Bear", names[f"out_{SID}_FV_Bear"], res.fair_values["Bear"])
    rows = _table_rows(xlsx, values, "VALUATION_OUTPUTS", "tbl_ValuationOutputs")
    row = next(r for r in rows if r["Security_ID"] == SID)
    assert blank(row["FV_Base_INR"]) and row["Valuation_Coverage_Status"] == "UNVALUED"


def test_terminal_growth_at_or_above_r_is_fail_and_never_a_number(built, data, tmp_path):
    aid = f"A-{SID}-Base-TERMINAL_GROWTH-ALL"
    xlsx = _edit(built.xlsx, tmp_path / OUTPUT_NAME, {("VALUATION_ASSUMPTIONS", _assumption_cell(built.xlsx, aid)): 0.13})
    values, names = built.recalc(xlsx)
    spec, res = _rerun(data, {aid: 0.13})
    assert res.terminal_flag_base == "FAIL" and res.fair_values["Base"] is None
    assert blank(names[f"out_{SID}_FV_Base"]) and blank(names[f"out_{SID}_TV_Share"])
    sheet = values[built.cell_map["sheet"]]
    assert sheet[built.cell("out.tflag")] == "FAIL"
    assert all(blank(sheet.get(built.cell(f"grid.rg.{i}.{j}"))) for i in range(5) for j in range(5))


def test_probabilities_and_reinvestment_controls(built, data, tmp_path):
    p = {"Bear": 0.25, "Base": 0.5, "Bull": 0.25}
    edits = {(built.cell_map["sheet"], built.cell(f"p.{c}")): x for c, x in p.items()}
    edits[(built.cell_map["sheet"], built.cell("r_reinv"))] = 0.07
    edits[(built.cell_map["sheet"], built.cell("er_mode"))] = REINVESTED
    xlsx = _edit(built.xlsx, tmp_path / OUTPUT_NAME, edits)
    values, _ = built.recalc(xlsx)
    _, res = _rerun(data, {}, reinvestment_mode=REINVESTED, reinvestment_rate=0.07, probabilities=p)
    sheet = values[built.cell_map["sheet"]]
    for c in P.CASES:
        er = res.cases[c].expected_return
        assert er.reinvested is not None and er.selected == er.reinvested
        assert_same(f"{c} reinvested", sheet[built.cell(f"{c}.reinv")], er.reinvested)
        assert_same(f"{c} selected", sheet[built.cell(f"{c}.er")], er.selected)
    assert sheet[built.cell("p_status")] == "PASS"
    assert_same("probability-weighted", sheet[built.cell("p_er")], res.probability_weighted_return)
    assert res.probability_weighted_return is not None and NOT_REINVESTED != REINVESTED


def test_probabilities_that_do_not_sum_to_one_fail(built, tmp_path):
    p = {"Bear": 0.3, "Base": 0.5, "Bull": 0.3}
    edits = {(built.cell_map["sheet"], built.cell(f"p.{c}")): x for c, x in p.items()}
    values, _ = built.recalc(_edit(built.xlsx, tmp_path / OUTPUT_NAME, edits))
    sheet = values[built.cell_map["sheet"]]
    assert sheet[built.cell("p_status")] == "FAIL" and blank(sheet.get(built.cell("p_er")))
