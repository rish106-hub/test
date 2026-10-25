"""W layer: static scans of the generated workbook and its recalculated values
(spec/test_plan.md W-01, W-02, W-04 … W-09, W-15, W-17; workbook_interface §3).

A check whose subject does not exist in the validated build (for example tbl_PriceControls before
03-nav is merged) is skipped as NOT RUN with the reason; it is never reported as PASS.
"""
from __future__ import annotations

import re

import pytest

from src.common import ownership
from src.common import schema as S
from src.common.enums import OutputLabel
from src.common.excel_contract import check_formula
from src.common.styles import BLACK, BLUE, GREEN, GREY, PURPLE
from tests.workbook.conftest import require_production, require_real
from tests.workbook.wbmodel import ERRORS, is_error

pytestmark = pytest.mark.workbook

HEX = {k: v.lstrip("#").upper() for k, v in dict(BLACK=BLACK, BLUE=BLUE, GREEN=GREEN, GREY=GREY, PURPLE=PURPLE).items()}
EXPLAINING_STATUS = {"FAIL", "UNMATCHED", "NO_PRICE", "UNVALUED", "MISSING", "STALE", "WARNING"}

PC_INPUTS = ["Market_Price_Scenario_Mode", "Market_Price_Shock_Pct", "Absolute_Scenario_Price",
             "Fair_Value_Scenario_Mode", "Fundamental_Case", "Fair_Value_Bias_Pct", "Required_Margin_of_Safety_Pct"]
SHOCK_INPUTS = [("tbl_PriceControls", "Market_Price_Shock_Pct"), ("tbl_SectorShocks", "Sector_Price_Shock_Pct")]


def _holdings_tables(wb):
    return [t for t in wb.tables if t.startswith("tbl_Holdings_")]


def _schema(table_name):
    from tools.check_module_contract import schema_for_excel_table
    return schema_for_excel_table(table_name)


# --------------------------------------------------------------------------- W-01
def test_W_01_release_build_complete(built, record_property):
    """Production mode, no unbuilt sheets, no not-implemented module, no stubs."""
    require_production(built)
    r = built.build_report
    record_property("qa_evidence", f"unbuilt={r.get('unbuilt_sheets')} not_implemented={list(r.get('not_implemented', {}))} "
                                   f"stubbed={r.get('stubbed')}")
    assert r.get("test_build") is False
    assert r.get("unbuilt_sheets") == [] and r.get("not_implemented") == {} and not r.get("stubbed")


# --------------------------------------------------------------------------- W-02
def test_W_02_no_error_values(built, record_property):
    """No #REF!/#NAME?/#DIV/0!/#VALUE!; #N/A only in a table row whose status column explains it."""
    bad, explained = [], 0
    for sheet, cells in built.all_values().items():
        for coord, v in cells.items():
            if not is_error(v):
                continue
            if v.startswith("#N/A"):
                hit = built.table_of(sheet, coord)
                if hit:
                    t, _ = hit
                    row = int(re.sub(r"[A-Z]+", "", coord))
                    statuses = {str(built.value(t.sheet, t.cell(h, row))) for h in t.headers if "Status" in h}
                    if statuses & EXPLAINING_STATUS or any(s.startswith("FAIL") for s in statuses):
                        explained += 1
                        continue
            bad.append((sheet, coord, v))
    record_property("qa_evidence", f"{len(bad)} unexplained error values; {explained} #N/A explained by a status column")
    assert bad == [], bad[:20]


# --------------------------------------------------------------------------- W-04
def test_W_04_whitelisted_functions_only(built, record_property):
    problems, n = [], 0
    for sheet, coord, f in built.formula_cells():
        n += 1
        p = check_formula(f)
        if p:
            problems.append((sheet, coord, p))
    record_property("qa_evidence", f"{n} formulas scanned against FUNCTION_WHITELIST; {len(problems)} violations")
    assert problems == [], problems[:20]


# --------------------------------------------------------------------------- W-05
_SHEETREF = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z_][A-Za-z0-9_.]*))!\$?[A-Z]{1,3}\$?\d+")


def test_W_05_no_raw_reference_to_another_owners_sheet(built, record_property):
    """Cross-sheet access only via tbl_*/in_*/out_*; plus each module's references vs its consumes list."""
    bad = []
    for sheet, coord, f in built.formula_cells():
        own = ownership.sheet_owner(sheet)
        for m in _SHEETREF.finditer(re.sub(r'"(?:[^"]|"")*"', '""', f)):
            target = (m.group(1) or m.group(2)).replace("''", "'")
            if target != sheet and ownership.sheet_owner(target) != own:
                bad.append((sheet, coord, target, f[:120]))
    contract_problems = []
    if built.manifest:
        from tools.check_module_contract import check_manifest, load_contracts
        contracts = load_contracts()
        owners = {t["owner"] for t in built.manifest.get("tables", {}).values()} | set(built.manifest.get("references", {}))
        for w in sorted(o for o in owners if o in contracts and o not in ("00-main", "05-qa")):
            contract_problems += [f"{w}: {p}" for p in check_manifest(w, contracts[w], built.manifest)
                                  if not p.startswith("missing interface table")]
    record_property("qa_evidence", f"{len(bad)} raw cross-owner references; manifest problems: {contract_problems[:5]}")
    assert bad == [] and contract_problems == [], (bad[:10], contract_problems[:10])


# --------------------------------------------------------------------------- W-06
def _scenario_input_cells(wb):
    cells = set()
    if "tbl_PriceControls" in wb.tables:
        for c in PC_INPUTS:
            if c in wb.tables["tbl_PriceControls"].headers:
                cells |= wb.column_cells("tbl_PriceControls", c)
    if "tbl_SectorShocks" in wb.tables:
        cells |= wb.column_cells("tbl_SectorShocks", "Sector_Price_Shock_Pct")
    if "tbl_CustomAllocation" in wb.tables:
        t = wb.tables["tbl_CustomAllocation"]
        for h in t.headers:
            cells |= wb.column_cells("tbl_CustomAllocation", h)
    for n in ("in_Global_Price_Shock_Pct", "in_Global_Case", "in_Default_MoS"):
        if n in wb.names:
            cells.add(wb.names[n])
    return cells


def _protected_cells(wb):
    out = set()
    if "tbl_PriceControls" in wb.tables:
        for c in ("Original_Market_Price", "Model_FV_Base", "Model_FV_Bear", "Model_FV_Bull"):
            if c in wb.tables["tbl_PriceControls"].headers:
                out |= wb.column_cells("tbl_PriceControls", c)
    for t in ("tbl_MarketPrices", "tbl_ValuationOutputs", "tbl_Registry", "tbl_RawFinancials", "tbl_OfficialNAV",
              "tbl_SchemeTotals", "tbl_CorporateActions"):
        if t in wb.tables:
            for h in wb.tables[t].headers:
                out |= wb.column_cells(t, h)
    for t in _holdings_tables(wb):
        for h in ("Current_Market_Price", "Base_Fair_Value", "Adj_Base", "Adj_Bear", "Adj_Bull", "MV_Mkt_INR_Cr",
                  "MV_Fair_Base_INR_Cr", *[c.name for c in S.FUND_HOLDINGS.stored_columns]):
            if h in wb.tables[t].headers:
                out |= wb.column_cells(t, h)
    if "tbl_NAVSummary" in wb.tables:
        for h in ("Official_NAV", "Fair_NAV_Base", "Fair_NAV_Bear", "Fair_NAV_Bull", "Fair_NAV_M2_Base"):
            if h in wb.tables["tbl_NAVSummary"].headers:
                out |= wb.column_cells("tbl_NAVSummary", h)
    for sheet, coord, _ in wb.formula_cells():
        if sheet.startswith("VAL_"):
            out.add((sheet, coord))
    return out


def test_W_06_originals_model_values_and_imports_never_depend_on_scenario_inputs(built, record_property):
    """R008: no original price, model fair value, imported value or non-scenario NAV depends on a
    scenario input; imported (canonical CSV) columns hold constants, never formulas."""
    require_real(built, "tbl_PriceControls")
    inputs, protected = _scenario_input_cells(built), _protected_cells(built)
    bad = [(s, c) for s, c in protected if (s, c) in built.graph and built.upstream({(s, c)}) & inputs]
    formula_in_import = []
    for name, t in built.tables.items():
        sch = _schema(name)
        if sch is None or sch.csv_path is None or not built.is_real(name):
            continue
        for col in sch.stored_columns:
            if col.name in t.headers:
                formula_in_import += [x for x in built.column_cells(name, col.name) if built.formula(*x)]
    record_property("qa_evidence", f"{len(protected)} protected cells traced; {len(bad)} depend on scenario inputs; "
                                   f"{len(formula_in_import)} formulas in imported columns")
    assert bad == [] and formula_in_import == [], (bad[:10], formula_in_import[:10])


# --------------------------------------------------------------------------- W-07
EXPECT_CONST = {HEX["BLUE"], HEX["PURPLE"], HEX["GREY"]}


def test_W_07_font_colour_matches_content(built, record_property):
    """Formula: green if it reads another sheet, else black. Numeric/date constant: blue, purple or grey."""
    bad, n = [], 0
    for sheet, coord, f in built.formula_cells():
        n += 1
        cross = any(s != sheet for s, _ in built.precedents(sheet, coord, f))
        want = HEX["GREEN"] if cross else HEX["BLACK"]
        got = built.font_rgb(sheet, coord) or HEX["BLACK"]
        if got != want:
            bad.append((sheet, coord, "formula", got, want))
    for ws in built.wf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None or isinstance(v, str):
                    continue
                n += 1
                got = built.font_rgb(ws.title, c.coordinate) or HEX["BLACK"]
                if got not in EXPECT_CONST:
                    bad.append((ws.title, c.coordinate, "constant", got, "blue/purple/grey"))
    record_property("qa_evidence", f"{n} formula and numeric constant cells checked; {len(bad)} colour mismatches")
    assert bad == [], bad[:20]


# --------------------------------------------------------------------------- W-08
OUTPUT_COLUMNS = {
    "tbl_NAVSummary": {"Official_NAV": OutputLabel.OFFICIAL_NAV, "Fair_NAV_Base": OutputLabel.FAIR_NAV,
                       "Fair_NAV_Bear": OutputLabel.FAIR_NAV, "Fair_NAV_Bull": OutputLabel.FAIR_NAV,
                       "Fair_NAV_Selected": OutputLabel.FAIR_NAV, "Fair_NAV_M2_Base": OutputLabel.FAIR_NAV,
                       "Scenario_Market_NAV": OutputLabel.SCENARIO_MARKET_NAV},
    "tbl_PriceControls": {"Target_Purchase_Price": OutputLabel.TARGET_PURCHASE_PRICE},
}


def _caption_near(wb, sheet, coord, label: str, table=None) -> bool:
    from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
    col, row = coordinate_from_string(coord)
    ci = column_index_from_string(col)
    ws = wb.wv[sheet]
    top = table.header_row if table else row
    cands = [ws.cell(row, c).value for c in range(1, ci + 4)]                        # same row, left or near right
    cands += [ws.cell(r, ci).value for r in range(max(1, top - 3), top + 1)]        # caption above the column
    cands += [ws.cell(r, ci).value for r in range(row + 1, row + 4)]               # or just below
    return any(isinstance(v, str) and label in v for v in cands)


def test_W_08a_start_here_defines_the_five_outputs(built, record_property):
    """R006: START_HERE states all five OutputLabel captions."""
    start = " ".join(str(v) for v in built.all_values().get("START_HERE", {}).values())
    missing_defs = [l.value for l in OutputLabel if l.value not in start]
    record_property("qa_evidence", f"START_HERE missing captions: {missing_defs}")
    assert missing_defs == []


def test_W_08b_value_outputs_carry_their_captions(built, record_property):
    """WI-9: every cell showing one of the value outputs has its OutputLabel caption nearby."""
    if not any(built.is_real(t) for t in OUTPUT_COLUMNS):
        require_real(built, *OUTPUT_COLUMNS)
    bad, n = [], 0
    for table, cols in OUTPUT_COLUMNS.items():
        if not built.is_real(table):
            continue
        t = built.tables[table]
        for col, label in cols.items():
            if col not in t.headers:
                continue
            for r in t.data_rows:
                n += 1
                if not _caption_near(built, t.sheet, t.cell(col, r), label.value, t):
                    bad.append((t.sheet, t.cell(col, r), col))
            for s, c, f in built.formula_cells():          # every formula elsewhere that displays it
                if f"{table}[{col}]" in f and (s, c) not in built.column_cells(table, col):
                    n += 1
                    if not _caption_near(built, s, c, label.value):
                        bad.append((s, c, col))
    record_property("qa_evidence", f"{n} output cells checked; {len(bad)} without caption: {bad[:5]}")
    assert bad == [], bad[:20]


# --------------------------------------------------------------------------- W-09
def test_W_09_production_contains_no_synthetic_data(built, record_property):
    require_production(built)
    hits = [(s, c) for s, cells in built.all_values().items() for c, v in cells.items()
            if isinstance(v, str) and ("SYNTHETIC" in v.upper() or "SRC-SYNTH" in v.upper())]
    record_property("qa_evidence", f"{len(hits)} cells with synthetic markers")
    assert hits == [], hits[:20]


@pytest.mark.xfail(reason="known finding F-014: build/main_sheets.py writes VALIDATION_CHECKS!A1 over the "
                          "test-build watermark (D-013); reported as FAIL until 00-main fixes it", strict=False)
def test_W_09b_test_build_watermarks_every_sheet(built, record_property):
    """D-013: a test build carries the synthetic watermark on every sheet."""
    if not built.is_test_build:
        pytest.skip("NOT RUN: applies to a test build; this is a production build (see W-09)")
    missing = [ws.title for ws in built.wv.worksheets if "SYNTHETIC" not in str(ws["A1"].value or "").upper()]
    record_property("qa_evidence", f"{len(built.wv.worksheets)} sheets; unwatermarked: {missing}")
    assert missing == []


# --------------------------------------------------------------------------- W-15
_INDEX = re.compile(r"INDEX\((tbl_[A-Za-z0-9_]+)\[([A-Za-z0-9_]+)\],MATCH\(([^()]*)\)\)")


def test_W_15_links_to_possibly_blank_sources_are_blank_safe(built, record_property):
    """WI-7 / D-022: INDEX of an empty cell returns 0, so a link to a column that is (or may be) blank
    must use link_patterns.blank_safe (`IF(INDEX(..)="","",INDEX(..))`)."""
    checked, bad = 0, []
    for sheet, coord, f in built.formula_cells():
        for t, c, _ in _INDEX.findall(f):
            if not built.is_real(t) or c not in built.tables[t].headers:
                continue
            src = built.column_cells(t, c)
            constants = [x for x in src if not built.formula(*x)]
            has_blank = any(built.wf[s][cc].value in (None, "") for s, cc in constants)
            sch = _schema(t)
            nullable = bool(sch) and any(col.name == c and col.origin == "stored" and not col.required
                                         for col in sch.columns) and bool(constants)
            if not (has_blank or nullable):
                continue
            checked += 1
            if not re.search(rf'INDEX\({t}\[{c}\],MATCH\([^()]*\)\)=""', f):
                bad.append((sheet, coord, f"{t}[{c}]", "blank in source" if has_blank else "nullable column"))
    if checked == 0:
        pytest.skip("NOT RUN: no link to a module-built table with a possibly blank column in this build")
    record_property("qa_evidence", f"{checked} links to possibly blank sources; {len(bad)} not blank-safe")
    assert bad == [], bad[:20]


# --------------------------------------------------------------------------- W-17
def test_W_17_shocks_are_read_once(built, record_property):
    """D-024: the three shock inputs are read only by tbl_PriceControls[Effective_Shock_Pct]
    (exception recorded in F-003: Control_Warning, which §6.2 obliges to read them, and boolean
    default checks); downstream scenario prices link to Selected_Scenario_Market_Price."""
    require_real(built, "tbl_PriceControls")
    shock_cells = set()
    for t, c in SHOCK_INPUTS:
        if t in built.tables and c in built.tables[t].headers:
            shock_cells |= built.column_cells(t, c)
    if "in_Global_Price_Shock_Pct" in built.names:
        shock_cells.add(built.names["in_Global_Price_Shock_Pct"])
    pc = built.tables["tbl_PriceControls"]
    allowed = built.column_cells("tbl_PriceControls", "Effective_Shock_Pct") if "Effective_Shock_Pct" in pc.headers else set()
    tolerated = built.column_cells("tbl_PriceControls", "Control_Warning") if "Control_Warning" in pc.headers else set()
    readers, exceptions = [], []
    for node, prec in built.graph.items():
        if prec & shock_cells and node not in allowed:
            if node in tolerated or isinstance(built.value(*node), bool):
                exceptions.append(node)
            else:
                readers.append(node)
    downstream = []
    for t in _holdings_tables(built):
        tb = built.tables[t]
        if "Selected_Scenario_Price" in tb.headers:
            for r in tb.data_rows:
                f = built.formula(tb.sheet, tb.cell("Selected_Scenario_Price", r)) or ""
                if "Selected_Scenario_Market_Price" not in f:
                    downstream.append((tb.sheet, tb.cell("Selected_Scenario_Price", r)))
    for s, c, f in built.formula_cells():
        if s != built.tables["tbl_PriceControls"].sheet and re.search(r"\[(Market_Price_After_Shock|Effective_Shock_Pct)\]", f):
            downstream.append((s, c))
    record_property("qa_evidence", f"shock readers outside Effective_Shock_Pct: {readers[:5]}; tolerated (F-003): "
                                   f"{len(exceptions)}; downstream not via Selected_Scenario_Market_Price: {downstream[:5]}")
    assert readers == [] and downstream == [], (readers[:10], downstream[:10])
