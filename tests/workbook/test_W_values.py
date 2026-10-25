"""W layer: recalculated values (spec/test_plan.md W-03, W-11 … W-14, W-16, W-18, W-19).

W-03 compares every row of tbl_PriceControls and tbl_NAVSummary with (a) the 05-qa clean-room
implementation and (b) the workers' reference functions, both fed the inputs read from the same
recalculated workbook, to 1e-9 relative. Nothing here trusts a cached value.
"""
from __future__ import annotations

import re
from datetime import date

import pytest

from src.common.enums import OutputLabel
from tests.independent._support import same
from tests.independent.cleanroom import controls as QC, nav as QN
from tests.workbook import wbread as RD
from tests.workbook.conftest import require_real
from tests.workbook.wbmodel import as_date, blankish, is_error, num

pytestmark = pytest.mark.workbook
REL = 1e-9


# =========================================================================== W-03 controls
def _pc_compare(wb, resolve, values=None) -> list:
    diffs = []
    for row in wb.rows("tbl_PriceControls", values):
        inp = RD.pc_inputs(wb, row, values)
        out = resolve(inp)
        for col, field in RD.PC_OUT.items():
            if col not in row:
                diffs.append((row["Security_ID"], col, "column missing"))
                continue
            got, want = RD.cell(row[col]), getattr(out, field)
            want = None if want == "" and field == "control_warning" else want
            if field == "fv_label" and str(inp["fair_value_mode"]) == QC.FV_MANUAL:
                ok = isinstance(got, str) and got.startswith("MANUALLY BIASED fair value (")   # F-006
            elif field == "control_status" and isinstance(want, str) and want.startswith("FAIL"):
                ok = isinstance(got, str) and got.startswith("FAIL")
            else:
                ok = same(got, want, REL)
            if not ok:
                diffs.append((row["Security_ID"], col, got, want))
    return diffs


def test_W_03a_price_controls_equal_cleanroom(built, record_property):
    require_real(built, "tbl_PriceControls")
    d = _pc_compare(built, lambda i: QC.resolve(QC.Controls(**{k: (str(v) if isinstance(v, str) else v) for k, v in i.items()})))
    record_property("qa_evidence", f"{len(built.rows('tbl_PriceControls'))} securities x {len(RD.PC_OUT)} outputs; "
                                   f"{len(d)} disagreements with the clean-room: {d[:3]}")
    assert d == [], d[:20]


def test_W_03b_price_controls_equal_worker_reference(built, record_property):
    require_real(built, "tbl_PriceControls")
    from src.scenarios.pricing import PriceControlInputs, resolve_controls
    from tests.independent._support import call_worker
    call_worker(resolve_controls, PriceControlInputs(750.0, 500.0))
    d = _pc_compare(built, lambda i: resolve_controls(PriceControlInputs(**i)))
    record_property("qa_evidence", f"{len(d)} disagreements with src.scenarios.pricing.resolve_controls: {d[:3]}")
    assert d == [], d[:20]


# =========================================================================== W-03 NAV summary
def _expected_nav_row(wb, row, fns, values=None) -> dict:
    m1, m2, m3, hcov, vcov, fresh = fns
    fund = row["Fund_ID"]
    lines = RD.lines_from_workbook(wb, fund, values)
    off = num(row["Official_NAV"])
    V = as_date(RD.global_input(wb, "in_Valuation_Date", values))
    pdt = RD.portfolio_date(wb, fund, values)
    reported = RD.scheme_net_assets(wb, fund, pdt, values)
    g = lambda n: num(RD.global_input(wb, n, values))   # noqa: E731
    e = {}
    for case, col in (("Base", "Fair_NAV_Base"), ("Bear", "Fair_NAV_Bear"), ("Bull", "Fair_NAV_Bull"),
                      ("Selected", "Fair_NAV_Selected")):
        e[col] = m1(off, lines, case).nav
    has_mpt = any(QN.coverage_status(l) == QN.MARKET_PASS_THROUGH for l in lines)
    if not has_mpt:                                            # F-001: basis ambiguous for MPT positions
        e["Fair_NAV_M2_Base"] = m2(off, lines, "Base").nav
        e["Scenario_Market_NAV"] = m3(off, lines).nav
        e["Method_Drift_INR"] = None if e["Fair_NAV_M2_Base"] is None else e["Fair_NAV_Base"] - e["Fair_NAV_M2_Base"]
    e["Diff_vs_Official_INR"] = e["Fair_NAV_Base"] - off
    e["Diff_vs_Official_Pct"] = e["Fair_NAV_Base"] / off - 1
    h = hcov(lines, reported, g("in_Recon_Tolerance_Pct"))
    e.update(Holdings_Value_Coverage_Pct=h.value_coverage_pct, Holdings_Match_Pct=h.match_pct,
             Unmatched_Count=len(h.unmatched_line_ids), Recon_Status=h.recon_status, Holdings_Coverage_Status=h.status,
             Recon_Gap_Pct=sum(l.market_value_inr_cr for l in lines) / reported - 1)
    c = vcov(lines, g("in_Coverage_Warn"), g("in_Coverage_Fail"))
    e.update(Valuation_Coverage_Total_Pct=c.coverage_total_pct, Valuation_Coverage_Equity_Pct=c.coverage_equity_pct,
             Unvalued_Weight_Pct=c.unvalued_weight_pct, Unvalued_Count=c.unvalued_count,
             Unvalued_Value_INR_Cr=c.unvalued_value_inr_cr, Valuation_Coverage_Status=c.status, Coverage_Caption=c.caption)
    f = fresh(lines, pdt, as_date(row["NAV_Date"]), V, int(g("in_Holdings_Stale_Days")),
              int(g("in_Financials_Stale_Days")), int(g("in_Price_Stale_Days")))
    e.update(Holdings_Age_Days=f.holdings_age_days, NAV_Date_Gap_Days=f.nav_date_gap_days,
             Max_Financials_Age_Days=f.max_financials_age_days, Stale_Valued_Weight_Pct=f.stale_valued_weight_pct,
             Freshness_Status=f.status)
    e["Portfolio_Date"] = pdt
    e["Valuation_Date"] = V
    e["Price_Date"] = as_date(RD.global_input(wb, "in_Price_Date", values))
    return e


def _official_nav_expected(wb, plan, values=None):
    """§7.1: NAV at in_Price_Date, else the latest earlier NAV."""
    pd_ = as_date(RD.global_input(wb, "in_Price_Date", values))
    obs = [(as_date(r["NAV_Date"]), num(r["NAV_INR"])) for r in wb.rows("tbl_OfficialNAV", values) if r["Plan_ID"] == plan]
    obs = [o for o in obs if o[0] and o[0] <= pd_]
    return max(obs)[1] if obs else None


def nav_summary_diffs(wb, fns, values=None) -> list:
    diffs = []
    for row in wb.rows("tbl_NAVSummary", values):
        want = _expected_nav_row(wb, row, fns, values)
        want["Official_NAV"] = _official_nav_expected(wb, row["Plan_ID"], values)
        for col, w in want.items():
            if col not in row:
                diffs.append((row["Plan_ID"], col, "column missing"))
                continue
            got = RD.cell(row[col])
            if isinstance(w, date):
                got = as_date(got)
            if not same(got, w, REL):
                diffs.append((row["Plan_ID"], col, got, w))
    return diffs


CLEANROOM_FNS = (QN.weighted_adjustment_nav, QN.quantity_based_nav, QN.scenario_market_nav, QN.holdings_coverage,
                 QN.coverage_summary, QN.freshness_summary)


def test_W_03c_nav_summary_equals_cleanroom(built, record_property):
    require_real(built, "tbl_NAVSummary", "tbl_PriceControls")
    d = nav_summary_diffs(built, CLEANROOM_FNS)
    record_property("qa_evidence", f"{len(built.rows('tbl_NAVSummary'))} plans; {len(d)} disagreements with the "
                                   f"clean-room: {d[:4]}")
    assert d == [], d[:20]


def test_W_03d_nav_summary_equals_worker_reference(built, record_property):
    require_real(built, "tbl_NAVSummary", "tbl_PriceControls")
    from src.fund_nav import coverage as C, methods as M
    from tests.contract import cases as K
    from tests.independent._support import call_worker
    call_worker(M.weighted_adjustment_nav, 80.0, K.MF001, "Base")
    d = nav_summary_diffs(built, (M.weighted_adjustment_nav, M.quantity_based_nav, M.scenario_market_nav,
                                  C.holdings_coverage, C.coverage_summary, C.freshness_summary))
    record_property("qa_evidence", f"{len(d)} disagreements with the 03-nav reference functions: {d[:4]}")
    assert d == [], d[:20]


def test_W_03e_out_names_equal_valuation_outputs(built, record_property):
    """Each VAL_* named output equals its tbl_ValuationOutputs row (the single seam to the NAV engine)."""
    require_real(built, "tbl_ValuationOutputs")
    vo = RD.by_key(built.rows("tbl_ValuationOutputs"), "Security_ID")
    pat = re.compile(r"^out_(S\d+)_(FV_Base|FV_Bear|FV_Bull|TV_Share)$")
    col = {"FV_Base": "FV_Base_INR", "FV_Bear": "FV_Bear_INR", "FV_Bull": "FV_Bull_INR", "TV_Share": "Terminal_Value_Share_Base"}
    diffs, n = [], 0
    for name in built.names:
        m = pat.match(name)
        if not m:
            continue
        n += 1
        sid, field = m.groups()
        got, want = RD.cell(built.name_value(name)), RD.cell(vo.get(sid, {}).get(col[field]))
        if not same(got, want, REL):
            diffs.append((name, got, want))
    if n == 0:
        pytest.skip("NOT RUN: no out_* names in this build")
    record_property("qa_evidence", f"{n} out_* names compared with tbl_ValuationOutputs; {len(diffs)} differ. "
                                   "Python re-derivation of out_* from VAL_* inputs needs F-012")
    assert diffs == [], diffs


# =========================================================================== W-11
def test_W_11_reset_identity_at_default_controls(built, record_property):
    """With every control at its default: Scenario Market NAV = Official NAV for every plan, and the
    workbook shows Controls_At_Default = TRUE (contract §6.3, §7.4)."""
    require_real(built, "tbl_PriceControls", "tbl_NAVSummary")
    mos = num(RD.global_input(built, "in_Default_MoS"))
    case = RD.global_input(built, "in_Global_Case")
    not_default = []
    for r in built.rows("tbl_PriceControls"):
        d = {"Market_Price_Scenario_Mode": "Original Market Price", "Fair_Value_Scenario_Mode": "Model Base Fair Value",
             "Fundamental_Case": case, "Required_Margin_of_Safety_Pct": mos}
        bad = [k for k, v in d.items() if not same(r.get(k), v)]
        bad += [k for k in ("Market_Price_Shock_Pct", "Fair_Value_Bias_Pct") if num(r.get(k)) not in (None, 0.0)]
        bad += ["Absolute_Scenario_Price"] if not blankish(r.get("Absolute_Scenario_Price")) else []
        if bad:
            not_default.append((r["Security_ID"], bad))
    if "tbl_SectorShocks" in built.tables:
        not_default += [("sector", r["Sector"]) for r in built.rows("tbl_SectorShocks") if num(r["Sector_Price_Shock_Pct"]) not in (None, 0.0)]
    if num(RD.global_input(built, "in_Global_Price_Shock_Pct")) not in (None, 0.0):
        not_default.append(("global", "in_Global_Price_Shock_Pct"))
    identity = [(r["Plan_ID"], r["Scenario_Market_NAV"], r["Official_NAV"]) for r in built.rows("tbl_NAVSummary")
                if not same(num(r["Scenario_Market_NAV"]), num(r["Official_NAV"]), 1e-12)]
    flags = [(s, c) for s, cells in built.all_values().items() for c, v in cells.items()
             if isinstance(v, str) and v.strip() == "Controls_At_Default"]
    flag_vals = []
    for s, c in flags:
        ws = built.wv[s]
        cc = ws[c]
        flag_vals += [ws.cell(cc.row, cc.column + 1).value, ws.cell(cc.row + 1, cc.column).value]
    found = any(v is True for v in flag_vals)
    record_property("qa_evidence", f"controls not at default: {not_default[:5]}; identity failures: {identity}; "
                                   f"Controls_At_Default cells: {flags} -> {flag_vals} (F-004)")
    assert not_default == [] and identity == [] and found, (not_default[:5], identity, flags)


# =========================================================================== W-12 / W-18 dashboard
def _dashboard_links(wb) -> dict[str, set]:
    """NAVSummary columns linked from EXECUTIVE_DASHBOARD -> the plan columns that show them."""
    out = {}
    for s, c, f in wb.formula_cells():
        if s == "EXECUTIVE_DASHBOARD":
            for col in re.findall(r"tbl_NAVSummary\[([A-Za-z0-9_]+)\]", f):
                if col != "Plan_ID":
                    out.setdefault(col, set()).add(re.sub(r"\d+", "", c))
    return out


def _plans(wb):
    return [r["Plan_ID"] for r in wb.rows("tbl_NAVSummary")]


def test_W_12_dashboard_shows_coverage_unvalued_and_date_warnings(built, record_property):
    require_real(built, "tbl_NAVSummary")
    need = ["Coverage_Caption", "Unvalued_Count", "Unvalued_Weight_Pct", "Holdings_Age_Days", "Price_Date",
            "NAV_Date_Gap_Days", "Freshness_Status"]
    links, n = _dashboard_links(built), len(_plans(built))
    missing = [c for c in need if len(links.get(c, ())) < n]
    shown_missing = [(s, c) for s, c, _ in built.formula_cells() if s == "EXECUTIVE_DASHBOARD"
                     and (built.value(s, c) == "MISSING" or is_error(built.value(s, c)))]
    record_property("qa_evidence", f"dashboard lacks per-plan {missing}; cells showing MISSING/error: {shown_missing[:5]} (F-011)")
    assert missing == [] and shown_missing == [], (missing, shown_missing[:10])


def test_W_18_three_quality_measures_separate_no_composite(built, record_property):
    require_real(built, "tbl_NAVSummary")
    links, n = _dashboard_links(built), len(_plans(built))
    statuses = ["Holdings_Coverage_Status", "Valuation_Coverage_Status", "Freshness_Status"]
    missing = [c for c in statuses if len(links.get(c, ())) < n]
    composite = [(s, c, v) for s, cells in built.all_values().items() for c, v in cells.items()
                 if isinstance(v, str) and re.search(r"(?i)\b(composite|overall (data )?quality|quality score|data score)\b", v)]
    record_property("qa_evidence", f"status links missing: {missing}; composite-score labels: {composite[:5]}")
    assert missing == [] and composite == [], (missing, composite)


# =========================================================================== W-13 charts
def test_W_13_charts_reference_cells_and_carry_metadata(built, record_property):
    zones = built.manifest.get("chart_zones", {})
    charts = [(ws.title, ch) for ws in built.wf.worksheets for ch in getattr(ws, "_charts", [])]
    if not zones and not charts:
        pytest.skip("NOT RUN: no chart zones or charts in this build (04-charts is Wave 2)")
    if built.build_report.get("not_implemented", {}).get("src.charts.render:render_charts"):
        pytest.skip("NOT RUN: src.charts.render:render_charts not implemented in this build (04-charts, Wave 2)")
    bad = []
    for sheet, ch in charts:
        for sr in ch.series:
            ref = getattr(getattr(sr, "val", None), "numRef", None)
            if ref is None or not getattr(ref, "f", None):
                bad.append((sheet, "series without a cell range"))
        if ch.title is None:
            bad.append((sheet, "chart without title"))
    from openpyxl.utils.cell import coordinate_from_string
    placed = {(s, ch.anchor._from.col, ch.anchor._from.row) for s, ch in charts}
    for zid, z in zones.items():
        col, row = coordinate_from_string(z["anchor"])
        from openpyxl.utils import column_index_from_string
        ci = column_index_from_string(col) - 1
        if (z["sheet"], ci, row - 1) not in placed:
            notice = built.value(z["sheet"], z["anchor"])
            if not (isinstance(notice, str) and re.search(r"(?i)missing|insufficient|not enough|no data", notice)):
                bad.append((zid, "zone has neither a chart nor a missing-data notice"))
    record_property("qa_evidence", f"{len(charts)} charts, {len(zones)} zones; problems: {bad[:5]}")
    assert bad == [], bad[:20]


# =========================================================================== W-14 terminal controls
def test_W_14a_tv_share_present_for_every_valued_security(built, record_property):
    require_real(built, "tbl_ValuationOutputs")
    bad = [(r["Security_ID"], r.get("Terminal_Value_Share_Base")) for r in built.rows("tbl_ValuationOutputs")
           if r.get("Valuation_Coverage_Status") == "MODEL_VALUED" and num(r.get("Terminal_Value_Share_Base")) is None]
    bad += [(r["Security_ID"], "FV present but UNVALUED") for r in built.rows("tbl_ValuationOutputs")
            if r.get("Valuation_Coverage_Status") == "UNVALUED" and any(num(r.get(c)) is not None
                                                                       for c in ("FV_Base_INR", "FV_Bear_INR", "FV_Bull_INR"))]
    record_property("qa_evidence", f"valued rows without TV share / unvalued rows with a FV: {bad}")
    assert bad == []


def test_W_14b_no_gordon_cell_with_g_at_or_above_r_shows_a_number(built):
    """TV-1/TV-5. The sensitivity grids on VAL_* sheets have no interface (table or names) from which
    QA can read r and g per cell, so this cannot be executed generically. Finding F-012."""
    require_real(built, "tbl_ValuationOutputs")
    pytest.skip("NOT RUN: VAL_* sensitivity-grid cells are not exposed through a tbl_*/out_* interface (F-012)")


# =========================================================================== W-16 units
def test_W_16_units_rederived(built, record_property):
    """MV_Mkt_INR_Cr = q × P / 1e7 and MV_Fair_Base_INR_Cr = q × FV / 1e7 per position; ONA lines carry
    the disclosed crore value; Close_INR = Close_LCY × FX."""
    require_real(built, "tbl_PriceControls")
    tables = [t for t in built.tables if t.startswith("tbl_Holdings_") and built.is_real(t)]
    if not tables:
        require_real(built, "tbl_Holdings_MF001")
    bad, n = [], 0
    for t in tables:
        for r in built.rows(t):
            q, p, fv = num(r.get("Quantity_Held")), num(r.get("Current_Market_Price")), num(r.get("Base_Fair_Value"))
            if r["Line_Role"] == "POSITION" and q is not None and p is not None:
                n += 1
                if not same(num(r.get("MV_Mkt_INR_Cr")), q * p / 1e7, REL):
                    bad.append((t, r["Holding_Line_ID"], "MV_Mkt", r.get("MV_Mkt_INR_Cr"), q * p / 1e7))
                want_fair = q * (fv if r.get("Valuation_Coverage_Status") == "MODEL_VALUED" else p) / 1e7
                if "MV_Fair_Base_INR_Cr" in r and not same(num(r.get("MV_Fair_Base_INR_Cr")), want_fair, REL):
                    bad.append((t, r["Holding_Line_ID"], "MV_Fair_Base", r.get("MV_Fair_Base_INR_Cr"), want_fair))
            elif r["Line_Role"] != "POSITION":
                n += 1
                if not same(num(r.get("MV_Mkt_INR_Cr")), num(r.get("Market_Value_INR_Cr")), REL):
                    bad.append((t, r["Holding_Line_ID"], "ONA MV", r.get("MV_Mkt_INR_Cr")))
    if built.is_real("tbl_MarketPrices"):
        for r in built.rows("tbl_MarketPrices"):
            n += 1
            if not same(num(r["Close_INR"]), num(r["Close_LCY"]) * num(r["FX_Rate_to_INR"]), 1e-6):
                bad.append(("tbl_MarketPrices", r["Security_ID"], r["Close_INR"]))
    record_property("qa_evidence", f"{n} unit re-derivations; {len(bad)} mismatches. Per-share BV from "
                                   "crore/shares needs VAL_* interface (F-012)")
    assert bad == [], bad[:20]


# =========================================================================== W-19 date captions
def test_W_19_reported_values_dated_and_estimates_captioned(built, record_property):
    """D-025: every display of BV_Reported_At_Anchor_INR says "as of <Anchor_Date>"; every value
    derived for V (BV_Estimated_At_Valuation_INR) is captioned "estimate"."""
    require_real(built, "tbl_ValuationOutputs")
    t = built.tables["tbl_ValuationOutputs"]
    bad, n = [], 0
    for col, needle in (("BV_Reported_At_Anchor_INR", "as of"), ("BV_Estimated_At_Valuation_INR", "estimate")):
        if col not in t.headers:
            bad.append((col, "column missing"))
            continue
        sources = set()
        for r in t.data_rows:
            f = built.formula(t.sheet, t.cell(col, r))
            if f:
                sources |= {p for p in built.precedents(t.sheet, t.cell(col, r), f) if p[0] != t.sheet}
        from openpyxl.utils.cell import coordinate_from_string
        for s, c in sources:                       # the VAL_* cell that displays the value
            n += 1
            _, row = coordinate_from_string(c)
            texts = " ".join(str(v) for v in [x.value for x in built.wv[s][row]] if v is not None).lower()
            if needle not in texts:
                bad.append((s, c, col, f"row lacks '{needle}'"))
    if n == 0:
        pytest.skip("NOT RUN: tbl_ValuationOutputs BV columns do not link to a VAL_* display in this build")
    record_property("qa_evidence", f"{n} VAL_* BV displays checked; {len(bad)} without date/estimate caption")
    assert bad == [], bad[:20]
